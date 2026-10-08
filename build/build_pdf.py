#!/usr/bin/env python3
"""Build the complete reader PDF from build/combined.md."""

from __future__ import annotations

import io
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parent.parent
INPUT = ROOT / "build" / "combined.md"
OUTPUT = ROOT / "output" / "kitzur-shulchan-aruch-questions-1-221.pdf"
FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
FONT_BOLD = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")


def clean_inline(s: str) -> str:
    return re.sub(r"(?<!\\)[*_\u0060]", "", s).strip()


def chapters_from_markdown(markdown: str):
    chapters = []
    current = None
    question = None
    for line in markdown.splitlines():
        m = re.match(r"^# Симан (\d+)\s*[—-]\s*(.+)$", line)
        if m:
            if current:
                if question:
                    current["items"].append(("question", question))
                chapters.append(current)
            current = {"number": int(m.group(1)), "title": clean_inline(m.group(2)), "items": []}
            question = None
            continue
        if not current:
            continue
        if line.startswith("## Вопросы") or line.strip() == "---":
            continue
        if line.startswith("### "):
            if question:
                current["items"].append(("question", question))
                question = None
            current["items"].append(("subheading", clean_inline(line[4:])))
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m:
            if question:
                current["items"].append(("question", question))
            question = f"{m.group(1)}. {clean_inline(m.group(2))}"
        elif line.strip() and question:
            question += " " + clean_inline(line)
    if current:
        if question:
            current["items"].append(("question", question))
        chapters.append(current)
    nums = [ch["number"] for ch in chapters]
    if nums != list(range(1, 222)):
        raise ValueError("Chapter sequence is not 1–221")
    return chapters


def body_docx(chapters, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Serif"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    heading = doc.styles["Heading 1"]
    heading.font.name = "DejaVu Sans"
    heading.font.size = Pt(17)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(29, 56, 80)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(16)
    heading.paragraph_format.keep_with_next = True

    sub = doc.styles["Heading 2"]
    sub.font.name = "DejaVu Sans"
    sub.font.size = Pt(10)
    sub.font.bold = True
    sub.font.color.rgb = RGBColor(75, 90, 103)
    sub.paragraph_format.space_before = Pt(9)
    sub.paragraph_format.space_after = Pt(5)
    sub.paragraph_format.keep_with_next = True

    for ch in chapters:
        p = doc.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = True
        p.add_run(f"Симан {ch['number']} — {ch['title']}")
        for kind, value in ch["items"]:
            if kind == "subheading":
                doc.add_paragraph(value, style="Heading 2")
            else:
                p = doc.add_paragraph(style="Normal")
                p.paragraph_format.left_indent = Cm(0.65)
                p.paragraph_format.first_line_indent = Cm(-0.65)
                p.paragraph_format.keep_together = True
                p.add_run(value)
    doc.save(path)


def registered_fonts() -> None:
    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT)))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_BOLD)))


def make_front(chapters, starts: dict[int, int], path: Path, toc_pages: int) -> None:
    c = Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFillColor(colors.HexColor("#17354b"))
    c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("DejaVu-Bold", 27)
    c.drawCentredString(w / 2, h - 210, "Кицур Шульхан Арух")
    c.setFont("DejaVu", 22)
    c.drawCentredString(w / 2, h - 260, "Проверочные вопросы")
    c.setFont("DejaVu", 17)
    c.drawCentredString(w / 2, h - 304, "Симаним 1–221")
    c.setStrokeColor(colors.HexColor("#b79d6f"))
    c.setLineWidth(2)
    c.line(110, h - 332, w - 110, h - 332)
    c.setFont("DejaVu", 10)
    c.drawCentredString(w / 2, 106, "На основе оригинального текста р. Шломо Ганцфрида")
    c.showPage()

    per_page = 30
    for page in range(toc_pages):
        c.setFillColor(colors.HexColor("#17354b"))
        c.setFont("DejaVu-Bold", 18)
        c.drawString(52, h - 62, "Оглавление" if page == 0 else "Оглавление (продолжение)")
        c.setStrokeColor(colors.HexColor("#b79d6f"))
        c.line(52, h - 74, w - 52, h - 74)
        for j, ch in enumerate(chapters[page * per_page : (page + 1) * per_page]):
            y = h - 108 - j * 23
            title = f"{ch['number']}. {ch['title']}"
            while pdfmetrics.stringWidth(title, "DejaVu", 9.5) > w - 145:
                title = title[:-2]
            if title != f"{ch['number']}. {ch['title']}":
                title = title.rstrip() + "…"
            c.setFont("DejaVu", 9.5)
            c.drawString(54, y, title)
            c.drawRightString(w - 54, y, str(starts[ch["number"]] + 1 + toc_pages))
        c.showPage()
    c.save()


def page_overlays(total: int) -> PdfReader:
    out = io.BytesIO()
    c = Canvas(out, pagesize=A4)
    w, h = A4
    for number in range(2, total + 1):
        c.setStrokeColor(colors.HexColor("#d5dce0"))
        c.line(50, 43, w - 50, 43)
        c.setFont("DejaVu", 8)
        c.setFillColor(colors.HexColor("#5b6870"))
        c.drawString(52, 28, "Кицур Шульхан Арух • Проверочные вопросы")
        c.drawRightString(w - 52, 28, f"{number} / {total}")
        c.showPage()
    c.save()
    out.seek(0)
    return PdfReader(out)


def main() -> None:
    registered_fonts()
    chapters = chapters_from_markdown(INPUT.read_text(encoding="utf-8"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ksa-pdf-") as tmp:
        tmp = Path(tmp)
        docx = tmp / "body.docx"
        body_docx(chapters, docx)
        subprocess.run(
            ["soffice", "-env:UserInstallation=file://" + str(tmp / "lo-profile"),
             "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(docx)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        )
        body = PdfReader(str(tmp / "body.pdf"))
        starts = {}
        for i, page in enumerate(body.pages):
            text = page.extract_text() or ""
            for m in re.finditer(r"Симан\s+(\d+)\s*[—-]", text):
                n = int(m.group(1))
                starts.setdefault(n, i + 1)
        if set(starts) != set(range(1, 222)):
            missing = sorted(set(range(1, 222)) - set(starts))
            raise ValueError(f"PDF headings missing: {missing}")
        toc_pages = (len(chapters) + 29) // 30
        front_path = tmp / "front.pdf"
        make_front(chapters, starts, front_path, toc_pages)
        front = PdfReader(str(front_path))
        writer = PdfWriter()
        writer.append(front)
        writer.append(body)
        total = len(writer.pages)
        overlays = page_overlays(total)
        for i in range(1, total):
            writer.pages[i].merge_page(overlays.pages[i - 1])
        for ch in chapters:
            writer.add_outline_item(
                f"Симан {ch['number']} — {ch['title']}",
                starts[ch["number"]] + len(front.pages) - 1,
            )
        with OUTPUT.open("wb") as stream:
            writer.write(stream)
        print(f"{OUTPUT}: {total} pages, {len(chapters)} chapters")


if __name__ == "__main__":
    main()
